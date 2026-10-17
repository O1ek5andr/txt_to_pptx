import io
import os
import re
import logging
from dotenv import load_dotenv
from telegram import Update
from telegram.ext import ApplicationBuilder, CommandHandler, MessageHandler, filters, ContextTypes
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

# Завантажуємо змінні оточення з .env файлу
# Абсолютний шлях до .env поруч із bot.py
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
ENV_PATH = os.path.join(BASE_DIR, ".env")

load_dotenv(dotenv_path=ENV_PATH)

# Налаштування логування
logging.basicConfig(
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
    level=logging.INFO
)

# Константи формату презентації (16:9)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FONT_FACE = "Arial Narrow"
MARGIN = 0.5

BASE_SIZE = 88
ALT_SIZE = 80

# --- ЛОГІКА ГЕНЕРАЦІЇ ПРЕЗЕНТАЦІЇ ---

def get_filename_from_first_line(raw_text: str) -> str:
    for row in raw_text.splitlines():
        line = row.strip()
        if not line or line.startswith("#") or (line.startswith("[") and line.endswith("]")):
            continue
        clean_name = re.sub(r'[\s.,;:!?\s]+$', '', line)
        return f"{clean_name}.pptx"
    return "Presentation.pptx"

def split_into_two_lines_perfectly(line: str):
    line = line.strip().upper()
    if len(line) <= 20:
        return [line]
    
    mid = len(line) // 2
    left_space = line.rfind(' ', 0, mid + 5)
    right_space = line.find(' ', mid - 5)
    
    if left_space == -1 and right_space == -1:
        return [line]
        
    split_point = left_space if left_space != -1 else right_space
    return [line[:split_point].strip(), line[split_point:].strip()]

def parse_song(raw_text: str):
    parts = {}
    current = "default"
    for row in raw_text.splitlines():
        line = row.strip()
        if not line or line.startswith("#"): continue
        m = re.match(r"^\[(.+?)\]$", line)
        if m:
            current = m.group(1).strip()
            parts.setdefault(current, [])
            continue
        parts.setdefault(current, []).append(line)
    return parts

def slide_needs_downsizing(slide_lines, size):
    usable_width_pts = (13.333 - (MARGIN * 2)) * 72
    for line in slide_lines:
        estimated_width = len(line) * (size * 0.58)
        if estimated_width > usable_width_pts:
            return True
    return False

def add_slide(prs, lines):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(0, 0, 0)
    
    title_shape = slide.shapes.title
    title_shape.left = Inches(MARGIN)
    title_shape.top = Inches(MARGIN)
    title_shape.width = prs.slide_width - Inches(MARGIN * 2)
    title_shape.height = prs.slide_height - Inches(MARGIN * 2)
    
    tf = title_shape.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = False

    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)

    font_size = ALT_SIZE if slide_needs_downsizing(lines, BASE_SIZE) else BASE_SIZE

    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.text = line.upper()
        
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        
        run = p.runs[0]
        run.font.name = FONT_FACE
        run.font.bold = True
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor(255, 255, 255)

def create_presentation_stream(raw_text: str) -> tuple[io.BytesIO, str]:
    parts = parse_song(raw_text)
    source_lines = []
    for p in parts.values():
        source_lines.extend(p)

    slides_content = [split_into_two_lines_perfectly(line) for line in source_lines]
    
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    
    for content in slides_content:
        add_slide(prs, content)
        
    # Збереження у буфер (в пам'ять)
    pptx_io = io.BytesIO()
    prs.save(pptx_io)
    pptx_io.seek(0)
    
    filename = get_filename_from_first_line(raw_text)
    return pptx_io, filename

# --- ОБРОБНИКИ ТЕЛЕГРАМ-БОТА ---

async def start_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "Привіт! Надішли мені текст пісні (повідомленням або .txt файлом), "
        "і я згенерую з нього готовий .pptx файл для презентації."
    )

async def handle_text_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_text = update.message.text
    if not user_text.strip():
        return

    await update.message.reply_text("Обробляю текст та створюю презентацію...")
    
    try:
        pptx_stream, filename = create_presentation_stream(user_text)
        await update.message.reply_document(
            document=pptx_stream,
            filename=filename,
            caption=""
        )
    except Exception as e:
        logging.error(f"Помилка створення презентації: {e}")
        await update.message.reply_text("Сталася помилка під час створення презентації.")

async def handle_document_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    document = update.message.document
    if not document.file_name.endswith('.txt'):
        await update.message.reply_text("Будь ласка, надішли текстовий файл у форматі .txt")
        return

    await update.message.reply_text("Завантажую та обробляю файл...")
    
    try:
        telegram_file = await context.bot.get_file(document.file_id)
        file_bytes = await telegram_file.download_as_bytearray()
        song_text = file_bytes.decode('utf-8')

        pptx_stream, filename = create_presentation_stream(song_text)
        await update.message.reply_document(
            document=pptx_stream,
            filename=filename,
            caption=""
        )
    except Exception as e:
        logging.error(f"Помилка обробки файлу: {e}")
        await update.message.reply_text("Не вдалося прочитати файл. Переконайся, що він у кодуванні UTF-8.")

if __name__ == "__main__":
    BOT_TOKEN = os.getenv("BOT_TOKEN")
    
    if not BOT_TOKEN:
        raise ValueError("Токен бота не знайдено! Встанови BOT_TOKEN у середовищі або у файлі .env.")

    app = ApplicationBuilder().token(BOT_TOKEN).build()

    app.add_handler(CommandHandler("start", start_command))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_text_message))
    app.add_handler(MessageHandler(filters.Document.ALL, handle_document_message))

    logging.info("Бот запущений.")
    app.run_polling()